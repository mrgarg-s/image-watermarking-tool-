from logging import root
import tkinter as tk
from tkinter import ttk, filedialog, colorchooser, messagebox, simpledialog
from PIL import Image, ImageTk, ImageDraw, ImageFont, ImageEnhance, PngImagePlugin
import os
import math
import io
import hashlib
import base64

# Drag and drop feature
try:
    from tkinterdnd2 import TkinterDnD, DND_FILES
    HAS_DND = True
except ImportError:
    HAS_DND = False

# Meta Encryption feature (AES)
try:
    from cryptography.fernet import Fernet
    HAS_CRYPTO = True
except ImportError:
    HAS_CRYPTO = False

# For Word Docx Export
try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# For Multi-Page PDF reading
try:
    import fitz  # PyMuPDF
    HAS_PDF = True
except ImportError:
    HAS_PDF = False


SUPPORTED_FORMATS = [
    ("All Supported Files", "*.png;*.jpg;*.jpeg;*.pdf"),
    ("PDF Documents", "*.pdf"),
    ("Images", "*.png;*.jpg;*.jpeg"),
    ("All Files", "*.*")
]

class VulcanProjectApp:
    def __init__(self, main_window):
        self.root = main_window
        self.root.title("VULCAN WATERMARKING TOOL")
        
        icon_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'icon.ico')
        if os.path.exists(icon_path):
            try:
                self.root.iconbitmap(icon_path)
            except Exception as e:
                print(f"Could not set icon: {e}")

        self.root.geometry("1200x800")
        self.root.minsize(900, 650)
        
        # Native UI Colors (Dark mode)
        self.dark_bg = "#121212"
        self.panel_bg = "#1e1e1e"
        self.text_color_ui = "#ffffff"
        self.purple_theme = "#BB86FC" 
        self.root.configure(bg=self.dark_bg)
        
        # --- App State Variables ---
        self.base_image = None
        self.logo_pic = None
        self.watermark_color = (255, 255, 255) 
        self.custom_font = "arial.ttf" 
        
        # Font list dictionary { "Display Name" : "File Path" }
        self.available_fonts = {"Arial (Default)": "arial.ttf"}
        
        # --- Multi-page State ---
        self.document_pages = [] 
        self.current_page_idx = 0
        self.page_states = {} 
        
        # Zoom and Pan
        self.zoom_lvl = 1.0
        self.pan_x = 0
        self.pan_y = 0
        self.start_px = 0
        self.start_py = 0
        
        # Optimization cache 
        self.fast_preview_img = None
        self.last_canvas_size = (0, 0)
        self.last_enhancements = (1.0, 1.0, 1.0)
        self.scale_down_ratio = 1.0
        
        # Enhancement sliders
        self.val_bright = tk.DoubleVar(value=1.0)
        self.val_contrast = tk.DoubleVar(value=1.0)
        self.val_sharp = tk.DoubleVar(value=1.0)
        
        # Watermark position
        self.text_pos_x, self.text_pos_y = 0.5, 0.5 
        self.logo_pos_x, self.logo_pos_y = 0.5, 0.5 
        self.old_mouse_x = 0
        self.old_mouse_y = 0
        
        # Tiling Variables
        self.tiled_var_t = tk.BooleanVar(value=False)
        self.tiled_var_l = tk.BooleanVar(value=False)
        
        # Bounding box logic
        self.is_clicked = False
        self.what_is_user_doing = 'nothing'
        self.snap_x = False
        self.snap_y = False
        self.t_bounds = (0, 0, 0, 0)
        self.l_bounds = (0, 0, 0, 0)
        self.canvas_img_x = 0
        self.canvas_img_y = 0
        self.preview_width = 1
        self.preview_height = 1

        self.rotate_start_angle = 0.0   
        self.rotate_base_val   = 0.0    
        self.resize_start_dist = 1.0    
        self.resize_base_val   = 1.0    
        self.bbox_center       = (0, 0) 
        
        self.percentage_labels = {}
        
        self.history_stack = []
        self.current_step = -1
        self.is_loading_history = False 

        self.make_ui()
        self.update_percentage_labels()

        self.root.after(100, lambda: self.main_paned.sash_place(0, 400, 0))

    def get_state_dict(self):
        return {
            't_px': self.text_pos_x, 't_py': self.text_pos_y, 
            'l_px': self.logo_pos_x, 'l_py': self.logo_pos_y, 
            'zoom': self.zoom_lvl, 'px': self.pan_x, 'py': self.pan_y, 
            'wm_text': self.text_box.get("1.0", "end-1c"), 'wm_color': self.watermark_color, 'font': self.custom_font,
            't_size': self.slider_size.get(), 't_opac': self.slider_opac.get(), 't_rot': self.slider_rot.get(),
            't_stroke': self.slider_stroke.get(),
            'l_scale': self.logo_s.get(), 'l_opac': self.logo_o.get(), 'l_rot': self.logo_r.get(),
            'b': self.val_bright.get(), 'c': self.val_contrast.get(), 's': self.val_sharp.get(),
            'font_disp': self.font_display_var.get(),
            'tiled_t': self.tiled_var_t.get(), 'tiled_l': self.tiled_var_l.get(),
            'space_t': self.slider_spacing_t.get(), 'space_l': self.slider_spacing_l.get(),
            'meta_auth': self.meta_author.get(), 'meta_copy': self.meta_copyright.get(),
            'meta_enc': self.meta_encrypt.get(),
            'meta_enc_file': getattr(self, 'meta_encrypt_file', tk.BooleanVar(value=False)).get()
        }

    def save_current_state(self):
        if self.is_loading_history: return 
        state = self.get_state_dict()
        self.history_stack = self.history_stack[:self.current_step + 1]
        self.history_stack.append(state)
        self.current_step += 1
        
        if len(self.document_pages) > 1 and self.pdf_page_mode.get() == "custom":
            self.page_states[self.current_page_idx] = state

    def do_undo(self, event=None):
        if self.current_step > 0:
            self.current_step -= 1
            self.restore_from_state(self.history_stack[self.current_step])

    def do_redo(self, event=None):
        if self.current_step < len(self.history_stack) - 1:
            self.current_step += 1
            self.restore_from_state(self.history_stack[self.current_step])

    def restore_from_state(self, state, trigger_refresh=True):
        self.is_loading_history = True
        
        self.text_pos_x = state.get('t_px', 0.5)
        self.text_pos_y = state.get('t_py', 0.5)
        self.logo_pos_x = state.get('l_px', 0.5)
        self.logo_pos_y = state.get('l_py', 0.5)
        self.zoom_lvl = state.get('zoom', 1.0)
        self.pan_x = state.get('px', 0)
        self.pan_y = state.get('py', 0)
        self.watermark_color = state.get('wm_color', (255,255,255))
        self.custom_font = state.get('font', 'arial.ttf')
        
        font_disp = state.get('font_disp', 'Arial (Default)')
        self.font_display_var.set(font_disp)
        
        # Ensure the restored font is in the combobox list
        if font_disp not in self.available_fonts and self.custom_font != 'arial.ttf':
            self.available_fonts[font_disp] = self.custom_font
            self.font_combo['values'] = list(self.available_fonts.keys())
        
        self.text_box.delete("1.0", tk.END)
        self.text_box.insert("1.0", state.get('wm_text', ''))
        
        self.slider_size.set(state.get('t_size', 80))
        self.slider_opac.set(state.get('t_opac', 180))
        self.slider_rot.set(state.get('t_rot', 0))
        self.slider_stroke.set(state.get('t_stroke', 0))
        
        self.logo_s.set(state.get('l_scale', 15))
        self.logo_o.set(state.get('l_opac', 200))
        self.logo_r.set(state.get('l_rot', 0))
        
        self.val_bright.set(state.get('b', 1.0))
        self.val_contrast.set(state.get('c', 1.0))
        self.val_sharp.set(state.get('s', 1.0))
        self.slider_b.set(self.val_bright.get())
        self.slider_c.set(self.val_contrast.get())
        self.slider_s.set(self.val_sharp.get())
        
        self.tiled_var_t.set(state.get('tiled_t', False))
        self.tiled_var_l.set(state.get('tiled_l', False))
        self.slider_spacing_t.set(state.get('space_t', 50))
        self.slider_spacing_l.set(state.get('space_l', 50))
        
        self.meta_author.set(state.get('meta_auth', ""))
        self.meta_copyright.set(state.get('meta_copy', ""))
        self.meta_encrypt.set(state.get('meta_enc', False))
        if hasattr(self, 'meta_encrypt_file'):
            self.meta_encrypt_file.set(state.get('meta_enc_file', False))
        
        self.update_percentage_labels()
        if trigger_refresh: self.refresh_image()
        self.is_loading_history = False

    def build_slider_widget(self, container, label_text, min_val, max_val, dict_key, start_val):
        f = tk.Frame(container, bg=self.panel_bg)
        f.pack(fill=tk.X, pady=(5, 0)) 
        
        ttk.Label(f, text=label_text).pack(side=tk.LEFT)
        perc_label = tk.StringVar(value="")
        self.percentage_labels[dict_key] = perc_label
        
        tk.Label(f, textvariable=perc_label, bg=self.panel_bg, fg=self.purple_theme, font=("Arial", 9, "bold")).pack(side=tk.RIGHT)
        s = ttk.Scale(container, from_=min_val, to=max_val, orient=tk.HORIZONTAL)
        s.config(command=lambda v: self.slider_moved_event())
        s.set(start_val)
        s.pack(fill=tk.X, pady=(2, 2))
        s.bind("<ButtonRelease-1>", lambda e: self.save_current_state())
        return s

    def make_ui(self):
        s = ttk.Style()
        s.theme_use('clam')
        s.configure("TFrame", background=self.panel_bg)
        s.configure("TLabel", background=self.panel_bg, foreground=self.text_color_ui)
        s.configure("TButton", background="#333333", foreground=self.text_color_ui, font=("Arial", 10, "bold"), padding=5)
        s.map("TButton", background=[("active", self.purple_theme)])
        s.configure("TNotebook", background=self.dark_bg)
        s.configure("TNotebook.Tab", background="#2b2b2b", foreground=self.text_color_ui, padding=[10, 5])
        s.map("TNotebook.Tab", background=[("selected", self.purple_theme)])
        s.configure("TCheckbutton", background=self.panel_bg, foreground=self.text_color_ui)
        s.configure("TRadiobutton", background=self.panel_bg, foreground=self.text_color_ui)
        s.configure("TCombobox", background="#333", foreground="white", fieldbackground="#333")

        self.main_paned = tk.PanedWindow(self.root, orient=tk.HORIZONTAL, sashwidth=6, sashrelief=tk.RAISED, bg="#444444")
        self.main_paned.pack(fill=tk.BOTH, expand=True)
        
        self.tool_frame_wrapper = tk.Frame(self.main_paned, bg=self.panel_bg)
        self.canvas_frame = tk.Frame(self.main_paned, bg=self.dark_bg)
        
        self.main_paned.add(self.tool_frame_wrapper, minsize=350)
        self.main_paned.add(self.canvas_frame, minsize=400)

        self.tool_frame = tk.Frame(self.tool_frame_wrapper, bg=self.panel_bg, padx=15, pady=15)
        self.tool_frame.pack(fill=tk.BOTH, expand=True)

        self.main_canvas = tk.Canvas(self.canvas_frame, bg="#000000", highlightthickness=0, cursor="crosshair")
        self.main_canvas.pack(fill=tk.BOTH, expand=True, padx=15, pady=15)
        txt = "[ DRAG & DROP  PNG · JPEG · PDF  HERE ]\n\n🖱️  Drag watermark to move\n🟢  Green circle = Rotate\n⬜  Corner square = Resize\n🔍  Scroll to Zoom  •  Right-drag to Pan"
        self.canvas_msg = self.main_canvas.create_text(450, 350, text=txt, fill="#666666", font=("Arial", 14, "bold"), justify="center")

        self.root.bind('<Control-z>', self.do_undo)
        self.root.bind('<Control-y>', self.do_redo)

        self.z_box = tk.Frame(self.canvas_frame, bg="#202020", padx=10, pady=5)
        self.z_box.place(relx=1.0, rely=0.0, anchor="ne", x=-25, y=25)
        self.zoom_display = tk.StringVar(value="100%")
        tk.Label(self.z_box, textvariable=self.zoom_display, bg="#202020", fg=self.purple_theme, font=("Arial", 11, "bold")).pack(side=tk.LEFT, padx=5)
        tk.Button(self.z_box, text="-", command=lambda: self.change_zoom(-0.25), bg="#202020", fg="white", bd=0).pack(side=tk.LEFT, padx=5)
        tk.Button(self.z_box, text="+", command=lambda: self.change_zoom(0.25), bg="#202020", fg="white", bd=0).pack(side=tk.LEFT, padx=5)
        tk.Button(self.z_box, text="Reset", command=self.reset_camera, bg="#333", fg="white", bd=0).pack(side=tk.LEFT, padx=5)

        if HAS_DND:
            self.main_canvas.drop_target_register(DND_FILES)
            self.main_canvas.dnd_bind('<<Drop>>', self.handle_file_drop)

        self.main_canvas.bind("<ButtonPress-1>", self.click_down)
        self.main_canvas.bind("<B1-Motion>", self.mouse_dragging)
        self.main_canvas.bind("<ButtonRelease-1>", self.click_up) 
        self.main_canvas.bind("<ButtonPress-3>", self.right_click_down)
        self.main_canvas.bind("<B3-Motion>", self.right_click_drag)
        self.main_canvas.bind("<Button-2>", lambda e: self.reset_camera()) 
        self.main_canvas.bind("<MouseWheel>", self.scrolled)
        self.main_canvas.bind("<Button-4>", self.scrolled) 
        self.main_canvas.bind("<Button-5>", self.scrolled) 

        top_frame = ttk.Frame(self.tool_frame)
        top_frame.pack(fill=tk.X, side=tk.TOP)
        
        tk.Label(top_frame, text="VULCAN WATERMARKER", font=("Arial", 16, "bold"), bg=self.panel_bg, fg=self.purple_theme).pack(pady=(0, 10), anchor=tk.W)
        ttk.Button(top_frame, text="Upload File (Image/PDF)", command=self.select_file).pack(fill=tk.X, pady=(0,5))
        
        ur_frame = ttk.Frame(top_frame)
        ur_frame.pack(fill=tk.X, pady=(0,5))
        ttk.Button(ur_frame, text="↩ Undo", command=self.do_undo).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0,2))
        ttk.Button(ur_frame, text="↪ Redo", command=self.do_redo).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(2,0))
        
        bf = tk.Frame(self.tool_frame, bg=self.panel_bg)
        bf.pack(fill=tk.X, side=tk.BOTTOM, pady=(10, 0))
        ttk.Button(bf, text="Save / Export", command=self.export_img).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 2))
        ttk.Button(bf, text="Clear Everything", command=self.clear_all).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(2, 0))

        self.nav_container = tk.Frame(self.tool_frame, bg=self.panel_bg)
        self.nav_container.pack(fill=tk.X, side=tk.BOTTOM, pady=5)
        
        self.page_nav_frame = tk.Frame(self.nav_container, bg=self.panel_bg)
        self.pdf_page_mode = tk.StringVar(value="all")
        mode_frame = tk.Frame(self.page_nav_frame, bg=self.panel_bg)
        mode_frame.pack(fill=tk.X, pady=(0, 5))
        ttk.Radiobutton(mode_frame, text="Apply to ALL", variable=self.pdf_page_mode, value="all").pack(side=tk.LEFT, expand=True)
        ttk.Radiobutton(mode_frame, text="Custom per page", variable=self.pdf_page_mode, value="custom").pack(side=tk.LEFT, expand=True)
        
        nav_btns = tk.Frame(self.page_nav_frame, bg=self.panel_bg)
        nav_btns.pack(fill=tk.X)
        self.page_label_var = tk.StringVar(value="Page 1 of 1")
        ttk.Button(nav_btns, text="◀ Prev", command=self.prev_page).pack(side=tk.LEFT, fill=tk.X, expand=True)
        tk.Label(nav_btns, textvariable=self.page_label_var, bg=self.panel_bg, fg=self.text_color_ui, justify="center").pack(side=tk.LEFT, padx=10)
        ttk.Button(nav_btns, text="Next ▶", command=self.next_page).pack(side=tk.LEFT, fill=tk.X, expand=True)

        self.tabs = ttk.Notebook(self.tool_frame)
        self.tabs.pack(fill=tk.BOTH, expand=True, side=tk.TOP)
        self.tabs.bind("<<NotebookTabChanged>>", lambda e: self.tab_switched())
        
        self.tab_t = ttk.Frame(self.tabs)
        self.tab_l = ttk.Frame(self.tabs)
        self.tab_e = ttk.Frame(self.tabs)
        self.tab_meta = ttk.Frame(self.tabs) 
        self.tab_b = ttk.Frame(self.tabs)
        
        self.tabs.add(self.tab_t, text="Text")
        self.tabs.add(self.tab_l, text="Logo")
        self.tabs.add(self.tab_e, text="Filters")
        self.tabs.add(self.tab_meta, text="Output")
        self.tabs.add(self.tab_b, text="Batch")

        self.make_text_controls()
        self.make_logo_controls()
        self.make_filter_controls()
        self.make_meta_controls()
        self.make_batch_controls()

    def tab_switched(self):
        self.refresh_image()
        self.save_current_state()

    def make_text_controls(self):
        ttk.Label(self.tab_t, text="Type Text (Multi-line):").pack(anchor=tk.W, pady=(10,2))
        
        self.text_box = tk.Text(self.tab_t, height=4, font=("Arial", 11), bg="#222222", fg="#FFFFFF", 
                                insertbackground="#FFFFFF", relief="flat", highlightthickness=1, 
                                highlightbackground="#444", highlightcolor=self.purple_theme)
        self.text_box.pack(fill=tk.X, pady=(0, 5))
        self.text_box.bind("<KeyRelease>", lambda e: self.refresh_image())
        self.text_box.bind("<FocusOut>", lambda e: self.save_current_state())

        # FONT SELECTION FRAME UPDATE
        ff = ttk.Frame(self.tab_t)
        ff.pack(fill=tk.X, pady=2)
        ttk.Label(ff, text="Font:").pack(side=tk.LEFT)
        self.font_display_var = tk.StringVar(value="Arial (Default)")
        
        self.font_combo = ttk.Combobox(ff, textvariable=self.font_display_var, state="readonly", width=16)
        self.font_combo['values'] = list(self.available_fonts.keys())
        self.font_combo.pack(side=tk.LEFT, padx=(5, 5))
        self.font_combo.bind("<<ComboboxSelected>>", self.font_selected)

        ttk.Button(ff, text="Remove", command=self.remove_font).pack(side=tk.RIGHT, padx=(2,0))
        ttk.Button(ff, text="Add Font", command=self.pick_font_file).pack(side=tk.RIGHT)

        tf = ttk.Frame(self.tab_t)
        tf.pack(fill=tk.X, pady=(10, 2))
        ttk.Checkbutton(tf, text="Tiled Pattern Mode", variable=self.tiled_var_t, command=self.slider_moved_event).pack(side=tk.LEFT)

        self.slider_spacing_t = self.build_slider_widget(self.tab_t, "Tile Spacing", 0, 500, "t_space", 50)
        self.slider_size = self.build_slider_widget(self.tab_t, "Size", 10, 400, "text_size", 80)
        self.slider_stroke = self.build_slider_widget(self.tab_t, "Text Outline (Stroke)", 0, 10, "text_stroke", 0) 
        self.slider_opac = self.build_slider_widget(self.tab_t, "Transparency", 0, 255, "text_opacity", 180)
        self.slider_rot = self.build_slider_widget(self.tab_t, "Rotate", 0, 360, "text_rotate", 0)

        cf = tk.Frame(self.tab_t, bg=self.panel_bg)
        cf.pack(fill=tk.X, pady=(5, 0))
        ttk.Label(cf, text="Color:").pack(side=tk.LEFT, padx=(0, 5))
        
        my_colors = [("#FFFFFF", (255, 255, 255)), ("#000000", (0, 0, 0)), ("#FF0000", (255, 0, 0)), ("#00FF00", (0, 255, 0)), ("#0000FF", (0, 0, 255))]
        for h_code, rgb_val in my_colors:
            swatch = tk.Label(cf, text="    ", bg=h_code, cursor="hand2", relief="solid", bd=1)
            swatch.pack(side=tk.LEFT, padx=3)
            swatch.bind("<Button-1>", lambda e, r=rgb_val: self.change_text_color(r))
            
        ttk.Button(cf, text="More...", command=self.open_color_window).pack(side=tk.LEFT, padx=10)

    # --- FONT MANAGEMENT METHODS ---
    def font_selected(self, event=None):
        font_name = self.font_display_var.get()
        if font_name in self.available_fonts:
            self.custom_font = self.available_fonts[font_name]
            self.refresh_image()
            self.save_current_state()

    def pick_font_file(self):
        f = filedialog.askopenfilename(filetypes=[("Font Files", "*.ttf;*.otf")])
        if f: 
            font_name = os.path.basename(f)
            # Add to font list if it doesn't already exist
            if font_name not in self.available_fonts:
                self.available_fonts[font_name] = f
                self.font_combo['values'] = list(self.available_fonts.keys())
                
            self.font_display_var.set(font_name)
            self.custom_font = f
            self.refresh_image()
            self.save_current_state()
            
    def remove_font(self):
        font_name = self.font_display_var.get()
        if font_name == "Arial (Default)":
            messagebox.showwarning("Warning", "The default Arial font cannot be removed.")
            return
            
        if font_name in self.available_fonts:
            # Remove from dictionary and combobox
            del self.available_fonts[font_name]
            self.font_combo['values'] = list(self.available_fonts.keys())
            
            # Revert to default font
            self.font_display_var.set("Arial (Default)")
            self.custom_font = self.available_fonts["Arial (Default)"]
            self.refresh_image()
            self.save_current_state()
            messagebox.showinfo("Success", f"Font '{font_name}' has been removed.")

    def make_logo_controls(self):
        lf_btns = tk.Frame(self.tab_l, bg=self.panel_bg)
        lf_btns.pack(fill=tk.X, pady=(10,15))
        ttk.Button(lf_btns, text="Upload Logo", command=self.pick_logo).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 2))
        ttk.Button(lf_btns, text="Remove Logo", command=self.clear_logo).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(2, 0))
        
        lf = tk.Frame(self.tab_l, bg=self.panel_bg)
        lf.pack(fill=tk.X, pady=5)
        ttk.Checkbutton(lf, text="Tiled Pattern Mode", variable=self.tiled_var_l, command=self.slider_moved_event).pack(side=tk.LEFT)

        self.slider_spacing_l = self.build_slider_widget(self.tab_l, "Tile Spacing", 0, 500, "l_space", 50)
        self.logo_s = self.build_slider_widget(self.tab_l, "Size", 2, 100, "logo_scale", 15)
        self.logo_o = self.build_slider_widget(self.tab_l, "Visibility", 0, 255, "logo_opacity", 200)
        self.logo_r = self.build_slider_widget(self.tab_l, "Angle", 0, 360, "logo_rotate", 0)

    def make_filter_controls(self):
        ttk.Label(self.tab_e, text="Image Adjustments", font=("Arial", 10, "bold")).pack(anchor=tk.W, pady=(10, 5))

        self.slider_b = self.build_slider_widget(self.tab_e, "Brightness", 0.1, 2.5, "enh_b", 1.0)
        self.slider_c = self.build_slider_widget(self.tab_e, "Contrast",   0.1, 2.5, "enh_c", 1.0)
        self.slider_s = self.build_slider_widget(self.tab_e, "Sharpness",  0.0, 3.0, "enh_s", 1.0)

        def _sync_filter_vars(*_):
            self.val_bright.set(self.slider_b.get())
            self.val_contrast.set(self.slider_c.get())
            self.val_sharp.set(self.slider_s.get())

        self.slider_b.config(command=lambda v: (_sync_filter_vars(), self.slider_moved_event()))
        self.slider_c.config(command=lambda v: (_sync_filter_vars(), self.slider_moved_event()))
        self.slider_s.config(command=lambda v: (_sync_filter_vars(), self.slider_moved_event()))

        def reset_filters():
            self.slider_b.set(1.0); self.slider_c.set(1.0); self.slider_s.set(1.0)
            _sync_filter_vars()
            self.slider_moved_event()
            self.save_current_state()

        ttk.Button(self.tab_e, text="Reset Filters to Default", command=reset_filters).pack(fill=tk.X, pady=(10, 0))

    def make_meta_controls(self):
        ttk.Label(self.tab_meta, text="Export Format:", font=("Arial", 10, "bold")).pack(anchor=tk.W, pady=(10,5))
        self.export_fmt = tk.StringVar(value="PNG Image")
        fmt_cb = ttk.Combobox(self.tab_meta, textvariable=self.export_fmt, state="readonly", 
                              values=["PNG Image", "JPEG Image", "PDF Document", "Word (.docx)"])
        fmt_cb.pack(fill=tk.X)

        self.meta_author = tk.StringVar(value="")
        ttk.Label(self.tab_meta, text="Author Name:").pack(anchor=tk.W, pady=(10,0))
        ttk.Entry(self.tab_meta, textvariable=self.meta_author).pack(fill=tk.X, pady=(0, 5))

        self.meta_copyright = tk.StringVar(value="")
        ttk.Label(self.tab_meta, text="Copyright Text:").pack(anchor=tk.W)
        ttk.Entry(self.tab_meta, textvariable=self.meta_copyright).pack(fill=tk.X, pady=(0, 10))
        
        self.meta_encrypt = tk.BooleanVar(value=False)
        ttk.Checkbutton(self.tab_meta, text="Encrypt Meta Text", variable=self.meta_encrypt, command=self.save_current_state).pack(anchor=tk.W, pady=5)
        
        self.meta_encrypt_file = tk.BooleanVar(value=False)
        ttk.Checkbutton(self.tab_meta, text="Encrypt Entire Output File", variable=self.meta_encrypt_file, command=self.save_current_state).pack(anchor=tk.W, pady=(0,5))
        
        ttk.Label(self.tab_meta, text="Decryption Password:").pack(anchor=tk.W, pady=(5,0))
        self.meta_pwd = tk.StringVar()
        
        pwd_frame = tk.Frame(self.tab_meta, bg=self.panel_bg)
        pwd_frame.pack(fill=tk.X)

        self.pwd_entry = ttk.Entry(pwd_frame, textvariable=self.meta_pwd, show="*")
        self.pwd_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)

        def toggle_password():
            if self.pwd_entry.cget("show") == "*":
                self.pwd_entry.config(show="")
                self.pwd_toggle_btn.config(text="🙈") 
            else:
                self.pwd_entry.config(show="*")
                self.pwd_toggle_btn.config(text="👁")  

        self.pwd_toggle_btn = tk.Button(
            pwd_frame, text="👁", command=toggle_password,
            bg=self.panel_bg, fg=self.text_color_ui, bd=0, 
            activebackground=self.panel_bg, activeforeground=self.purple_theme,
            cursor="hand2"
        )
        self.pwd_toggle_btn.pack(side=tk.RIGHT, padx=(5, 0))

    def make_batch_controls(self):
        ttk.Label(self.tab_b, text="Batch Watermarking", font=("Arial", 11, "bold")).pack(anchor=tk.W, pady=(10, 2))
        ttk.Label(self.tab_b, text="Applies current watermark settings to\nmultiple images at once.").pack(anchor=tk.W, pady=(0, 10))

        fmt_frame = tk.Frame(self.tab_b, bg=self.panel_bg)
        fmt_frame.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(fmt_frame, text="Output Format:").pack(side=tk.LEFT, padx=(0, 8))
        self.batch_fmt = tk.StringVar(value="PNG Image")
        batch_fmt_cb = ttk.Combobox(fmt_frame, textvariable=self.batch_fmt, state="readonly", width=14,
                                    values=["PNG Image", "JPEG Image", "PDF Document", "Word (.docx)"])
        batch_fmt_cb.pack(side=tk.LEFT)

        ttk.Button(self.tab_b, text="Option 1: Select Specific Images", command=self.batch_specific_files).pack(fill=tk.X, pady=(0, 5))
        ttk.Button(self.tab_b, text="Option 2: Process Entire Folder",  command=self.batch_entire_folder).pack(fill=tk.X)

        self.batch_progress_var = tk.StringVar(value="")
        self.batch_progress_label = tk.Label(
            self.tab_b, textvariable=self.batch_progress_var,
            bg=self.panel_bg, fg=self.purple_theme, font=("Arial", 9), justify=tk.LEFT, wraplength=300
        )
        self.batch_progress_label.pack(anchor=tk.W, pady=(10, 0))

    def change_zoom(self, amount):
        if not self.base_image: return
        self.zoom_lvl = max(0.1, min(5.0, self.zoom_lvl + amount))
        self.zoom_display.set(f"{int(self.zoom_lvl * 100)}%")
        self.refresh_image()

    def scrolled(self, event):
        if event.num == 5 or event.delta < 0: val = -0.1
        else: val = 0.1
        self.change_zoom(val)

    def right_click_down(self, event):
        self.start_px = event.x
        self.start_py = event.y

    def right_click_drag(self, event):
        self.pan_x += (event.x - self.start_px)
        self.pan_y += (event.y - self.start_py)
        self.start_px = event.x
        self.start_py = event.y
        self.refresh_image()

    def reset_camera(self):
        self.pan_x = 0
        self.pan_y = 0
        self.zoom_lvl = 1.0
        self.zoom_display.set("100%")
        self.refresh_image()

    def slider_moved_event(self):
        if not self.is_loading_history:
            self.update_percentage_labels()
            self.refresh_image()
        
    def update_percentage_labels(self):
        try:
            self.percentage_labels["text_size"].set(f"{int(self.slider_size.get())}px")
            self.percentage_labels["text_stroke"].set(f"{int(self.slider_stroke.get())}px")
            self.percentage_labels["text_opacity"].set(f"{int((self.slider_opac.get()/255)*100)}%")
            self.percentage_labels["text_rotate"].set(f"{int(self.slider_rot.get())}°")
            self.percentage_labels["t_space"].set(f"{int(self.slider_spacing_t.get())}px")
            
            self.percentage_labels["logo_scale"].set(f"{int(self.logo_s.get())}%")
            self.percentage_labels["logo_opacity"].set(f"{int((self.logo_o.get()/255)*100)}%")
            self.percentage_labels["logo_rotate"].set(f"{int(self.logo_r.get())}°")
            self.percentage_labels["l_space"].set(f"{int(self.slider_spacing_l.get())}px")
            
            self.percentage_labels["enh_b"].set(f"{int(self.val_bright.get()*100)}%")
            self.percentage_labels["enh_c"].set(f"{int(self.val_contrast.get()*100)}%")
            self.percentage_labels["enh_s"].set(f"{int(self.val_sharp.get()*100)}%")
        except Exception as e:
            pass

    def get_real_bbox(self):
        if not self.base_image: return None
        try:
            tab = self.tabs.index("current")
        except:
            tab = 0

        if tab == 1: 
            if not self.logo_pic or self.l_bounds[2] == 0: return None
            orig_x, orig_y, orig_w, orig_h = self.l_bounds
        else: 
            if not self.text_box.get("1.0", "end-1c").strip() or self.t_bounds[2] == 0: return None
            orig_x, orig_y, orig_w, orig_h = self.t_bounds
            
        x = self.canvas_img_x + orig_x
        y = self.canvas_img_y + orig_y
        return x, y, orig_w, orig_h

    def paint_overlay(self):
        self.main_canvas.delete("ui_box")
        if not self.is_clicked: return
        b = self.get_real_bbox()
        if not b: return
        bx, by, bw, bh = b

        self.main_canvas.create_rectangle(
            bx, by, bx + bw, by + bh,
            outline=self.purple_theme, dash=(4, 4), width=2, tags="ui_box"
        )

        HANDLE_R = 8
        rx = bx + bw // 2
        ry = by - 18   
        self.main_canvas.create_line(
            bx + bw // 2, by, rx, ry + HANDLE_R,
            fill="#00cc66", width=2, tags="ui_box"
        )
        self.main_canvas.create_oval(
            rx - HANDLE_R, ry - HANDLE_R, rx + HANDLE_R, ry + HANDLE_R,
            fill="#00cc66", outline="#ffffff", width=2, tags="ui_box"
        )
        self.main_canvas.create_text(
            rx, ry, text="↻", fill="#ffffff", font=("Arial", 9, "bold"), tags="ui_box"
        )

        SQ = 9
        sx = bx + bw
        sy = by + bh
        self.main_canvas.create_rectangle(
            sx - SQ, sy - SQ, sx + SQ, sy + SQ,
            fill="#ffffff", outline=self.purple_theme, width=2, tags="ui_box"
        )

    def click_down(self, event):
        if not self.base_image: return
        self.old_mouse_x = event.x
        self.old_mouse_y = event.y
        b = self.get_real_bbox()
        if not b:
            self.what_is_user_doing = 'nothing'
            self.is_clicked = False
            self.paint_overlay()
            return

        bx, by, bw, bh = b
        cx = bx + bw // 2
        cy = by + bh // 2
        self.bbox_center = (cx, cy)

        SQ = 18
        if abs(event.x - (bx + bw)) <= SQ and abs(event.y - (by + bh)) <= SQ:
            self.what_is_user_doing = 'resizing'
            self.is_clicked = True
            self.resize_start_dist = max(1.0, math.hypot(event.x - cx, event.y - cy))
            is_text = (self.tabs.index("current") != 1)
            self.resize_base_val = self.slider_size.get() if is_text else self.logo_s.get()

        elif math.hypot(event.x - cx, event.y - (by - 18)) <= 18:
            self.what_is_user_doing = 'rotating'
            self.is_clicked = True
            self.rotate_start_angle = math.degrees(math.atan2(event.y - cy, event.x - cx))
            is_text = (self.tabs.index("current") != 1)
            self.rotate_base_val = self.slider_rot.get() if is_text else self.logo_r.get()

        elif bx <= event.x <= bx + bw and by <= event.y <= by + bh:
            self.what_is_user_doing = 'moving'
            self.is_clicked = True

        else:
            self.what_is_user_doing = 'nothing'
            self.is_clicked = False

        self.paint_overlay()

    def mouse_dragging(self, event):
        if not self.base_image or self.what_is_user_doing == 'nothing': return
        dx = event.x - self.old_mouse_x
        dy = event.y - self.old_mouse_y
        self.old_mouse_x = event.x
        self.old_mouse_y = event.y

        is_text = (self.tabs.index("current") != 1)
        cx, cy  = self.bbox_center

        if self.what_is_user_doing == 'moving':
            pw = max(1, self.preview_width)
            ph = max(1, self.preview_height)
            if is_text:
                self.text_pos_x = max(0.0, min(1.0, self.text_pos_x + dx / pw))
                self.text_pos_y = max(0.0, min(1.0, self.text_pos_y + dy / ph))
            else:
                self.logo_pos_x = max(0.0, min(1.0, self.logo_pos_x + dx / pw))
                self.logo_pos_y = max(0.0, min(1.0, self.logo_pos_y + dy / ph))

        elif self.what_is_user_doing == 'rotating':
            cur_angle  = math.degrees(math.atan2(event.y - cy, event.x - cx))
            delta      = cur_angle - self.rotate_start_angle
            new_val    = (self.rotate_base_val + delta) % 360
            if new_val < 0: new_val += 360
            if is_text:
                self.slider_rot.set(new_val)
            else:
                self.logo_r.set(new_val)
            self.update_percentage_labels()

        elif self.what_is_user_doing == 'resizing':
            cur_dist = max(1.0, math.hypot(event.x - cx, event.y - cy))
            ratio    = cur_dist / self.resize_start_dist
            if is_text:
                new_val = max(10, min(400, self.resize_base_val * ratio))
                self.slider_size.set(new_val)
            else:
                new_val = max(2, min(100, self.resize_base_val * ratio))
                self.logo_s.set(new_val)
            self.update_percentage_labels()

        self.refresh_image()

    def click_up(self, event):
        if self.what_is_user_doing != 'nothing': self.save_current_state()
        self.what_is_user_doing = 'nothing'
        self.paint_overlay()

    def process_file_input(self, filepath):
        if filepath.lower().endswith('.pdf'):
            if not HAS_PDF:
                messagebox.showerror("Dependency Error", "You need PyMuPDF to read PDFs.\nRun: pip install PyMuPDF")
                return
            
            try:
                doc = fitz.open(filepath)
                self.document_pages = []
                for i in range(len(doc)):
                    page = doc.load_page(i)
                    pix = page.get_pixmap(dpi=150) 
                    mode = "RGBA" if pix.alpha else "RGB"
                    img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
                    if mode != "RGBA": img = img.convert("RGBA")
                    self.document_pages.append(img)
                
                self.page_nav_frame.pack(fill=tk.X, pady=5) 
                self.export_fmt.set("PDF Document")
                
                self.history_stack = []
                self.current_step = -1
                self.page_states = {} 
                self.reset_camera()
                
                self.current_page_idx = 0
                self.load_page(0)
                self.save_current_state()
                
            except Exception as e:
                messagebox.showerror("PDF Error", f"Failed to load PDF: {e}")
        else:
            self.document_pages = []
            self.page_nav_frame.pack_forget() 
            img = Image.open(filepath).convert("RGBA")
            self.load_new_file(img)

    def load_page(self, idx):
        if not self.document_pages: return

        if self.base_image is not None:
            self.page_states[self.current_page_idx] = self.get_state_dict()

        if hasattr(self, 'pdf_page_mode') and self.pdf_page_mode.get() == "custom":
            if idx in self.page_states:
                self.restore_from_state(self.page_states[idx], trigger_refresh=False)
                
        self.current_page_idx = idx
        self.page_label_var.set(f"Page {idx+1} of {len(self.document_pages)}")
        self.base_image = self.document_pages[idx]
        self.fast_preview_img = None
        self.refresh_image()

    def next_page(self):
        if self.current_page_idx < len(self.document_pages) - 1:
            self.load_page(self.current_page_idx + 1)

    def prev_page(self):
        if self.current_page_idx > 0:
            self.load_page(self.current_page_idx - 1)

    def load_new_file(self, img_obj):
        self.base_image = img_obj
        self.fast_preview_img = None 
        self.reset_camera()
        self.history_stack = []
        self.current_step = -1
        self.page_states = {}
        self.refresh_image()
        self.save_current_state() 

    def select_file(self):
        f = filedialog.askopenfilename(filetypes=SUPPORTED_FORMATS)
        if f: self.process_file_input(f)

    def handle_file_drop(self, event):
        f = event.data
        if f.startswith('{') and f.endswith('}'): f = f[1:-1]
        self.process_file_input(f)

    def pick_logo(self):
        f = filedialog.askopenfilename(filetypes=[("Images", "*.png *.jpg *.jpeg")])
        if f:
            self.logo_pic = Image.open(f).convert("RGBA")
            self.refresh_image()
            self.save_current_state()
            
    def clear_logo(self):
        self.logo_pic = None
        self.refresh_image()
        self.save_current_state()

    def change_text_color(self, rgb_tuple):
        self.watermark_color = rgb_tuple
        self.refresh_image()
        self.save_current_state()

    def open_color_window(self):
        c = colorchooser.askcolor()
        if c[0]: 
            self.watermark_color = tuple(int(x) for x in c[0])
            self.refresh_image()
            self.save_current_state()

    def render_engine(self, img_source, low_res_mode=False):
        if img_source is None: return None
        
        if low_res_mode:
            cw = int(max(10, self.main_canvas.winfo_width()) * self.zoom_lvl)
            ch = int(max(10, self.main_canvas.winfo_height()) * self.zoom_lvl)
            temp = img_source.copy()
            temp.thumbnail((cw, ch), Image.Resampling.BILINEAR)
            self.scale_down_ratio = temp.width / img_source.width
            self.preview_width = temp.width
            self.preview_height = temp.height
            work_img = temp
            ratio = self.scale_down_ratio
        else:
            work_img = img_source.copy()
            ratio = 1.0

        w, h = work_img.size

        b_val = self.val_bright.get()
        c_val = self.val_contrast.get()
        s_val = self.val_sharp.get()
        if b_val != 1.0 or c_val != 1.0 or s_val != 1.0:
            alpha = work_img.getchannel("A") if work_img.mode == "RGBA" else None
            rgb   = work_img.convert("RGB")
            if b_val != 1.0: rgb = ImageEnhance.Brightness(rgb).enhance(b_val)
            if c_val != 1.0: rgb = ImageEnhance.Contrast(rgb).enhance(c_val)
            if s_val != 1.0: rgb = ImageEnhance.Sharpness(rgb).enhance(s_val)
            work_img = rgb.convert("RGBA")
            if alpha is not None:
                work_img.putalpha(alpha)   

        if low_res_mode:
            self.t_bounds = (0, 0, 0, 0)
            self.l_bounds = (0, 0, 0, 0)

        txt = self.text_box.get("1.0", "end-1c").strip()
        if txt:
            sz = max(1, int(self.slider_size.get() * ratio))
            stroke_width = int(self.slider_stroke.get() * ratio)
            
            try: 
                font = ImageFont.truetype(self.custom_font, sz)
            except: 
                font = ImageFont.load_default()
                stroke_width = 0 
            
            draw_dummy = ImageDraw.Draw(Image.new("RGBA", (1,1)))
            try:
                bbox = draw_dummy.multiline_textbbox((0,0), txt, font=font, stroke_width=stroke_width)
            except AttributeError:
                bbox = draw_dummy.multiline_textsize(txt, font=font, stroke_width=stroke_width)
                bbox = (0, 0, bbox[0], bbox[1])
            
            tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
            p = max(1, int(20 * ratio))
            t_img = Image.new('RGBA', (tw + p*2, th + p*2), (255, 255, 255, 0))
            draw = ImageDraw.Draw(t_img)
            
            r, g, b = self.watermark_color
            stroke_fill = (0,0,0) if (r*0.299 + g*0.587 + b*0.114) > 186 else (255,255,255)
            
            draw.multiline_text(
                (p - bbox[0], p - bbox[1]), txt, 
                fill=(*self.watermark_color, int(self.slider_opac.get())), 
                font=font, 
                stroke_width=stroke_width, 
                stroke_fill=(*stroke_fill, int(self.slider_opac.get())),
                align="center"
            )
            
            rot = self.slider_rot.get()
            if rot != 0: t_img = t_img.rotate(-rot, expand=True, resample=Image.Resampling.BILINEAR)
            nw, nh = t_img.size
            fx = int((w - nw) * self.text_pos_x)
            fy = int((h - nh) * self.text_pos_y)
            
            if self.tiled_var_t.get():
                sp = max(0, int(self.slider_spacing_t.get() * ratio))
                step_x = nw + sp
                step_y = nh + sp
                if step_x == 0: step_x = 1
                if step_y == 0: step_y = 1
                offset_x = fx % step_x
                offset_y = fy % step_y
                
                for cy in range(offset_y - step_y, h, step_y):
                    for cx in range(offset_x - step_x, w, step_x):
                        work_img.paste(t_img, (cx, cy), t_img)
            else:
                work_img.paste(t_img, (fx, fy), t_img)
            
            if low_res_mode: self.t_bounds = (fx, fy, nw, nh)

        if self.logo_pic:
            s_fact = self.logo_s.get() / 100.0
            md = min(w, h) * s_fact
            lw, lh = self.logo_pic.size
            r_scale = min(md/lw, md/lh)
            ns = (max(1, int(lw * r_scale)), max(1, int(lh * r_scale)))
            lr = self.logo_pic.resize(ns, Image.Resampling.BILINEAR)
            
            al = lr.split()[3]
            al = ImageEnhance.Brightness(al).enhance(self.logo_o.get() / 255.0)
            lr.putalpha(al)
            
            rot = self.logo_r.get()
            if rot != 0: lr = lr.rotate(-rot, expand=True, resample=Image.Resampling.BILINEAR)
            nw, nh = lr.size
            fx = int((w - nw) * self.logo_pos_x)
            fy = int((h - nh) * self.logo_pos_y)
            
            if self.tiled_var_l.get():
                sp = max(0, int(self.slider_spacing_l.get() * ratio))
                step_x = nw + sp
                step_y = nh + sp
                if step_x == 0: step_x = 1
                if step_y == 0: step_y = 1
                offset_x = fx % step_x
                offset_y = fy % step_y
                
                for cy in range(offset_y - step_y, h, step_y):
                    for cx in range(offset_x - step_x, w, step_x):
                        work_img.paste(lr, (cx, cy), lr)
            else:
                work_img.paste(lr, (fx, fy), lr)
            
            if low_res_mode: self.l_bounds = (fx, fy, nw, nh)

        return work_img

    def refresh_image(self):
        if self.base_image is None: return
        prev = self.render_engine(self.base_image, low_res_mode=True)
        self.display_image = ImageTk.PhotoImage(prev)
        self.main_canvas.delete("all")
        cw = max(10, self.main_canvas.winfo_width())
        ch = max(10, self.main_canvas.winfo_height())
        self.canvas_img_x = ((cw - prev.width) // 2) + self.pan_x
        self.canvas_img_y = ((ch - prev.height) // 2) + self.pan_y
        self.main_canvas.create_image(self.canvas_img_x, self.canvas_img_y, anchor=tk.NW, image=self.display_image)
        self.paint_overlay()

    def _encrypt_file_on_disk(self, filepath):
        if not self.meta_encrypt_file.get() or not HAS_CRYPTO or not self.meta_pwd.get():
            return filepath
        try:
            pwd_bytes = self.meta_pwd.get().encode()
            key = base64.urlsafe_b64encode(hashlib.sha256(pwd_bytes).digest())
            f = Fernet(key)
            with open(filepath, 'rb') as file:
                original = file.read()
            encrypted = f.encrypt(original)
            enc_filepath = filepath + ".enc"
            with open(enc_filepath, 'wb') as file:
                file.write(encrypted)
            os.remove(filepath) 
            return enc_filepath
        except Exception as e:
            print("Full File Encryption Error:", e)
            return filepath

    def _apply_meta_and_save(self, final_img, save_path, fmt):
        auth_txt = self.meta_author.get()
        copy_txt = self.meta_copyright.get()

        if self.meta_encrypt.get() and HAS_CRYPTO and self.meta_pwd.get():
            try:
                pwd_bytes = self.meta_pwd.get().encode()
                key = base64.urlsafe_b64encode(hashlib.sha256(pwd_bytes).digest())
                f = Fernet(key)
                auth_txt = f.encrypt(auth_txt.encode()).decode()
                copy_txt = f.encrypt(copy_txt.encode()).decode()
            except Exception as e:
                print("Encryption Error:", e)

        if fmt == "Word (.docx)":
            if not HAS_DOCX:
                messagebox.showerror("Error", "python-docx missing. Saved as PNG instead.")
                final_img.save(save_path.replace(".docx", ".png"), format="PNG")
            else:
                buf = io.BytesIO()
                final_img.save(buf, format="PNG")
                buf.seek(0)
                doc = docx.Document()
                doc.add_paragraph(f"Author: {auth_txt}\nCopyright: {copy_txt}")
                doc.add_picture(buf, width=docx.shared.Inches(6))
                doc.save(save_path)
        elif fmt == "PDF Document":
            final_img.convert("RGB").save(save_path, format="PDF", dpi=(150, 150))
        elif fmt == "JPEG Image":
            final_img.convert("RGB").save(save_path, format="JPEG", quality=95)
        else:  
            meta_info = PngImagePlugin.PngInfo()
            meta_info.add_text("Author", auth_txt)
            meta_info.add_text("Copyright", copy_txt)
            final_img.save(save_path, format="PNG", pnginfo=meta_info)

    def export_img(self):
        if self.base_image is None: return
        
        fmt = self.export_fmt.get()
        ext = ".png"
        if fmt == "JPEG Image":    ext = ".jpg"
        elif fmt == "PDF Document": ext = ".pdf"
        elif fmt == "Word (.docx)": ext = ".docx"

        save_path = filedialog.asksaveasfilename(defaultextension=ext)
        if not save_path: return
        
        if not save_path.lower().endswith(ext):
            save_path += ext
            
        if self.document_pages and fmt == "PDF Document":
            processed_pages = []
            original_ui_state = self.get_state_dict()

            for i, p in enumerate(self.document_pages):
                if hasattr(self, 'pdf_page_mode') and self.pdf_page_mode.get() == "custom":
                    state_for_page = self.page_states.get(i, original_ui_state)
                    self.restore_from_state(state_for_page, trigger_refresh=False)

                rendered = self.render_engine(p, low_res_mode=False)
                processed_pages.append(rendered.convert("RGB"))

            self.restore_from_state(original_ui_state, trigger_refresh=True)

            if processed_pages:
                try:
                    processed_pages[0].save(
                        save_path,
                        format="PDF",          
                        save_all=True,
                        append_images=processed_pages[1:],
                        dpi=(150, 150)         
                    )
                    final_path = self._encrypt_file_on_disk(save_path)
                    messagebox.showinfo("Done", f"Saved {len(processed_pages)}-page watermarked PDF!\nSaved to: {final_path}")
                except Exception as e:
                    messagebox.showerror("PDF Export Error", f"Could not save PDF:\n{e}")
            return

        final = self.render_engine(self.base_image, low_res_mode=False)
        self._apply_meta_and_save(final, save_path, fmt)
        final_path = self._encrypt_file_on_disk(save_path)
        messagebox.showinfo("Done", f"Saved successfully to:\n{final_path}")

    def execute_batch_loop(self, files_to_process):
        if not files_to_process: return
        
        o_dir = filedialog.askdirectory(title="Select Output Folder to save results")
        if not o_dir: return

        fmt = self.batch_fmt.get() if hasattr(self, 'batch_fmt') else self.export_fmt.get()

        ext_map = {
            "PNG Image":    ".png",
            "JPEG Image":   ".jpg",
            "PDF Document": ".pdf",
            "Word (.docx)": ".docx",
        }
        ext = ext_map.get(fmt, ".png")

        total   = len(files_to_process)
        count   = 0
        skipped = []

        self.batch_progress_var.set(f"Starting… 0 / {total}")
        self.root.update()

        for idx, path in enumerate(files_to_process, start=1):
            fname            = os.path.basename(path)
            name_without_ext = os.path.splitext(fname)[0]
            self.batch_progress_var.set(f"Processing {idx} / {total}: {fname}")
            self.root.update()

            try:
                lower = path.lower()
                if lower.endswith('.pdf'):
                    if not HAS_PDF:
                        skipped.append(f"{fname} (PyMuPDF not installed)")
                        continue
                    doc = fitz.open(path)
                    for p_idx in range(len(doc)):
                        page = doc.load_page(p_idx)
                        pix  = page.get_pixmap(dpi=150)
                        mode = "RGBA" if pix.alpha else "RGB"
                        p_img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
                        if mode != "RGBA":
                            p_img = p_img.convert("RGBA")
                        wm = self.render_engine(p_img, low_res_mode=False)
                        s_path = os.path.join(o_dir, f"vulcan_{name_without_ext}_p{p_idx+1}{ext}")
                        self._apply_meta_and_save(wm, s_path, fmt)
                        self._encrypt_file_on_disk(s_path)
                    count += 1
                else:
                    img = Image.open(path).convert("RGBA")
                    wm  = self.render_engine(img, low_res_mode=False)
                    s_path = os.path.join(o_dir, f"vulcan_{name_without_ext}{ext}")
                    self._apply_meta_and_save(wm, s_path, fmt)
                    self._encrypt_file_on_disk(s_path)
                    count += 1

            except Exception as e:
                skipped.append(f"{fname}: {e}")

        summary = f"Done! Processed {count} / {total} files."
        if skipped:
            summary += f"\nSkipped {len(skipped)} file(s)."
        self.batch_progress_var.set(summary)

        detail = f"Batch complete!\n\n✅ Processed: {count}\n❌ Skipped: {len(skipped)}"
        if skipped:
            detail += "\n\nSkipped files:\n" + "\n".join(f"• {s}" for s in skipped[:10])
            if len(skipped) > 10:
                detail += f"\n…and {len(skipped)-10} more."
        messagebox.showinfo("Batch Done", detail)

    def batch_specific_files(self):
        files = filedialog.askopenfilenames(
            title="Select Images or PDFs to Watermark",
            filetypes=[
                ("All Supported Files", "*.png *.jpg *.jpeg *.pdf"),
                ("PNG",  "*.png"),
                ("JPEG", "*.jpg *.jpeg"),
                ("PDF",  "*.pdf"),
                ("All Files", "*.*"),
            ]
        )
        if files: self.execute_batch_loop(files)

    def batch_entire_folder(self):
        in_dir = filedialog.askdirectory(title="Select Input Folder")
        if not in_dir: return

        BATCH_EXTS = ('.png', '.jpg', '.jpeg', '.pdf')
        files = [
            os.path.join(in_dir, f)
            for f in os.listdir(in_dir)
            if f.lower().endswith(BATCH_EXTS)
        ]

        if not files:
            messagebox.showwarning("Empty", "No supported images or PDFs found in that folder!")
            return

        self.execute_batch_loop(files)

    def clear_all(self):
        self.base_image = None
        self.logo_pic = None
        self.document_pages = []
        self.page_nav_frame.pack_forget()
        self.main_canvas.delete("all")

        self.main_canvas.create_text(
            450, 350,
            text="[ DRAG & DROP  PNG · JPEG · PDF  HERE ]\n\n🖱️  Drag watermark to move\n🟢  Green circle = Rotate\n⬜  Corner square = Resize\n🔍  Scroll to Zoom  •  Right-drag to Pan",
            fill="#666666", font=("Arial", 14, "bold"), justify="center"
        )

        self.text_box.delete("1.0", tk.END)
        self.meta_author.set("")
        self.meta_copyright.set("")

        self.history_stack = []
        self.current_step = -1
        self.page_states = {}

        self.reset_camera()

if __name__ == "__main__":
    if HAS_DND: 
        win = TkinterDnD.Tk()
    else: 
        win = tk.Tk()
        
    app = VulcanProjectApp(win)
    win.bind("<Configure>", lambda e: app.refresh_image() if e.widget == win else None)
    win.mainloop()